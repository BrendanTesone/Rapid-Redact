"""Excel export operations for search results and terms."""

import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont


def save_results_to_excel(
    results: list[dict[str, str | int]], output_path: str
) -> bool:
    """Export search results to Excel with formatting and highlighting."""
    if not results:
        return False

    df = pd.DataFrame(results)
    cols = ["file_name", "page", "term", "match", "context"]
    existing_cols = [c for c in cols if c in df.columns]
    df = df[existing_cols]
    df.to_excel(output_path, index=False)

    wb = load_workbook(output_path)
    ws = wb.active
    if ws is None:
        return False

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="0078D4", end_color="0078D4", fill_type="solid"
    )
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    context_col_idx = (
        int(df.columns.get_loc("context")) + 1 if "context" in df.columns else 5  # type: ignore[arg-type]
    )
    match_col_idx = (
        int(df.columns.get_loc("match")) + 1 if "match" in df.columns else 4  # type: ignore[arg-type]
    )

    for row in range(2, ws.max_row + 1):
        context_cell = ws.cell(row=row, column=context_col_idx)
        match_val = ws.cell(row=row, column=match_col_idx).value
        context_cell.alignment = Alignment(wrap_text=True, vertical="top")

        original_text = context_cell.value
        if original_text and match_val and str(match_val) in str(original_text):
            parts = str(original_text).partition(str(match_val))
            rich_string = CellRichText(
                parts[0],
                TextBlock(font=InlineFont(b=True, color="FF0000"), text=parts[1]),  # type: ignore[arg-type]
                parts[2],
            )
            context_cell.value = rich_string

    for col in ws.columns:
        col_letter = col[0].column_letter  # type: ignore[union-attr]
        ws.column_dimensions[col_letter].width = 60 if col_letter == "E" else 25

    wb.save(output_path)
    return True


def export_terms_to_excel(terms: list[dict[str, str | bool]], output_path: str) -> bool:
    """Export search terms to Excel with formatting."""
    wb = Workbook()
    ws = wb.active
    if ws is None:
        return False
    ws.title = "Search Terms"

    ws["A1"] = "Term"
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill(
        start_color="0078D4", end_color="0078D4", fill_type="solid"
    )
    ws["A1"].alignment = Alignment(horizontal="center")

    for i, term_item in enumerate(terms, start=2):
        term_value = (
            term_item.get("term", "")
            if isinstance(term_item, dict)
            else str(term_item)
        )
        ws.cell(row=i, column=1, value=term_value)

    max_length = max((len(str(t.get("term", ""))) for t in terms), default=10)
    ws.column_dimensions["A"].width = min(max(max_length + 2, 15), 50)

    if len(terms) > 0:
        ws.auto_filter.ref = f"A1:A{len(terms) + 1}"

    wb.save(output_path)
    return True


def import_terms_from_excel(file_path: str) -> list[str]:
    """Import search terms from Excel file."""
    wb = load_workbook(file_path, data_only=True)
    ws = wb.active
    if ws is None:
        return []

    first_cell_val = ws.cell(row=1, column=1).value
    is_header_row = isinstance(
        first_cell_val, str
    ) and first_cell_val.lower().strip() in {"term", "terms", "search term"}
    start_row = 2 if is_header_row else 1

    terms: list[str] = []
    for row in ws.iter_rows(
        min_row=start_row, max_row=ws.max_row, min_col=1, max_col=1
    ):
        cell_value = row[0].value
        if cell_value:
            term_str = str(cell_value).strip()
            if term_str:
                terms.append(term_str)

    return terms
